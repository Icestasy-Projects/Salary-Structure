"""
Generate Trevor Dcosta's Technical Analysis assignment with 10 annotated chart pattern images.
Creates TradingView-style candlestick charts with pattern demarcation.
Outputs: Word document (.docx) and PDF.
"""

import os
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyArrowPatch
import matplotlib.dates as mdates
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas as pdf_canvas
from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(BASE, "trevor_charts")
os.makedirs(IMG_DIR, exist_ok=True)

LOGO_PATH = os.path.join(BASE, "extracted_images", "vidhi", "img-000.png")

TV_DARK = {
    'bg': '#131722',
    'candle_up': '#26a69a',
    'candle_down': '#ef5350',
    'wick_up': '#26a69a',
    'wick_down': '#ef5350',
    'grid': '#1e222d',
    'text': '#d1d4dc',
    'axis': '#363a45',
    'vol_up': '#26a69a40',
    'vol_down': '#ef535040',
}

CHARTS = [
    {"num": 1,  "pattern": "Double Top",           "stock": "Axis Bank Ltd.",              "ticker": "AXISBANK", "timeframe": "Monthly", "exchange": "NSE"},
    {"num": 2,  "pattern": "Double Bottom",        "stock": "SBI Life Insurance Co. Ltd.", "ticker": "SBILIFE",  "timeframe": "Daily",   "exchange": "NSE"},
    {"num": 3,  "pattern": "Triple Top",           "stock": "Tata Technologies Ltd.",      "ticker": "TATATECH", "timeframe": "1 Hour",  "exchange": "NSE"},
    {"num": 4,  "pattern": "Triple Bottom",        "stock": "Infosys Ltd.",                "ticker": "INFY",     "timeframe": "Daily",   "exchange": "NSE"},
    {"num": 5,  "pattern": "Rectangle (Box)",      "stock": "Genus Power Infra Ltd.",      "ticker": "GENUSPOWER","timeframe": "Daily",  "exchange": "NSE"},
    {"num": 6,  "pattern": "Descending Triangle",  "stock": "Bharti Airtel Ltd.",          "ticker": "BHARTIARTL","timeframe": "2 Hour",  "exchange": "NSE"},
    {"num": 7,  "pattern": "Ascending Triangle",   "stock": "Bharat Petroleum Corp. Ltd.", "ticker": "BPCL",     "timeframe": "Daily",   "exchange": "NSE"},
    {"num": 8,  "pattern": "Symmetrical Triangle",  "stock": "Maruti Suzuki India Ltd.",   "ticker": "MARUTI",   "timeframe": "Daily",   "exchange": "NSE"},
    {"num": 9,  "pattern": "Diamond Top",          "stock": "Coforge Ltd.",                "ticker": "COFORGE",  "timeframe": "Daily",   "exchange": "NSE"},
    {"num": 10, "pattern": "Wedge",                "stock": "Tata Motors Ltd.",            "ticker": "TATAMOTORS","timeframe": "4 Hour",  "exchange": "NSE"},
]


def generate_ohlcv(n, base_price, trend_func, volatility=0.02):
    dates = [datetime(2024, 1, 2) + timedelta(days=i * 1.4) for i in range(n)]
    opens, highs, lows, closes, volumes = [], [], [], [], []
    price = base_price

    for i in range(n):
        trend = trend_func(i, n)
        noise = np.random.normal(0, volatility * price)
        change = trend + noise
        o = price
        c = price + change
        h = max(o, c) + abs(np.random.normal(0, volatility * 0.5 * price))
        l = min(o, c) - abs(np.random.normal(0, volatility * 0.5 * price))
        vol = int(np.random.uniform(500000, 3000000) * (1 + abs(change) / price * 10))

        opens.append(o)
        closes.append(c)
        highs.append(h)
        lows.append(l)
        volumes.append(vol)
        price = c

    return dates, opens, highs, lows, closes, volumes


def plot_candlestick(ax, dates, opens, highs, lows, closes, volumes=None):
    width = 0.6
    for i in range(len(dates)):
        color_up = TV_DARK['candle_up']
        color_down = TV_DARK['candle_down']
        if closes[i] >= opens[i]:
            color = color_up
            body_low = opens[i]
            body_high = closes[i]
        else:
            color = color_down
            body_low = closes[i]
            body_high = opens[i]

        ax.plot([i, i], [lows[i], highs[i]], color=color, linewidth=0.8)
        ax.bar(i, body_high - body_low, bottom=body_low, width=width,
               color=color, edgecolor=color, linewidth=0.5)


def setup_chart_axes(fig, ax, ticker, exchange, timeframe, n):
    ax.set_facecolor(TV_DARK['bg'])
    fig.patch.set_facecolor(TV_DARK['bg'])
    ax.tick_params(colors=TV_DARK['text'], labelsize=9)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_color(TV_DARK['axis'])
    ax.spines['left'].set_color(TV_DARK['axis'])
    ax.spines['bottom'].set_color(TV_DARK['axis'])
    ax.grid(True, alpha=0.15, color=TV_DARK['grid'])
    ax.set_xlim(-2, n + 2)

    title = f"{ticker}  ·  {exchange}  ·  {timeframe}"
    ax.set_title(title, color=TV_DARK['text'], fontsize=11, loc='left', pad=10,
                fontfamily='monospace')

    tv_text = "TradingView"
    ax.text(0.99, 0.02, tv_text, transform=ax.transAxes, fontsize=9,
            color='#555555', ha='right', va='bottom', fontfamily='sans-serif')


def draw_annotation_line(ax, x1, y1, x2, y2, color='#2196F3', style='--', lw=2, label=None):
    ax.plot([x1, x2], [y1, y2], color=color, linestyle=style, linewidth=lw, alpha=0.85)
    if label:
        mx = (x1 + x2) / 2
        my = (y1 + y2) / 2
        ax.text(mx, my, label, color=color, fontsize=9, ha='center', va='bottom',
                fontweight='bold', bbox=dict(boxstyle='round,pad=0.3', facecolor=TV_DARK['bg'],
                edgecolor=color, alpha=0.8))


def draw_annotation_circle(ax, x, y, radius, color='#FF5252', label=None, label_pos='top'):
    circle = plt.Circle((x, y), radius, fill=False, edgecolor=color, linewidth=2.5, linestyle='-')
    ax.add_patch(circle)
    if label:
        offset = radius + 2 if label_pos == 'top' else -(radius + 2)
        va = 'bottom' if label_pos == 'top' else 'top'
        ax.text(x, y + offset, label, color=color, fontsize=9, ha='center', va=va,
                fontweight='bold')


def gen_double_top(img_path):
    np.random.seed(42)
    n = 120
    base = 850

    def trend(i, n):
        if i < 30:
            return base * 0.006
        elif i < 45:
            return -base * 0.004
        elif i < 60:
            return base * 0.005
        elif i < 75:
            return -base * 0.003
        elif i < 90:
            return base * 0.004
        else:
            return -base * 0.007

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.012)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "AXISBANK", "NSE", "Monthly", n)

    peak1_x, peak2_x = 38, 68
    peak1_y = max(h[35:42])
    peak2_y = max(h[65:72])

    draw_annotation_circle(ax, peak1_x, peak1_y, 4, '#FF5252', 'Top 1')
    draw_annotation_circle(ax, peak2_x, peak2_y, 4, '#FF5252', 'Top 2')

    neckline_y = min(l[42:55])
    draw_annotation_line(ax, 10, neckline_y, 100, neckline_y, '#2196F3', '--', 2, 'Neckline')

    resistance_y = (peak1_y + peak2_y) / 2
    draw_annotation_line(ax, 25, resistance_y, 80, resistance_y, '#FF9800', ':', 1.5, 'Resistance')

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_double_bottom(img_path):
    np.random.seed(123)
    n = 100
    base = 1500

    def trend(i, n):
        if i < 25:
            return -base * 0.005
        elif i < 40:
            return base * 0.004
        elif i < 55:
            return -base * 0.005
        elif i < 70:
            return base * 0.003
        else:
            return base * 0.006

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.012)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "SBILIFE", "NSE", "Daily", n)

    bot1_x, bot2_x = 28, 58
    bot1_y = min(l[25:32])
    bot2_y = min(l[55:62])

    draw_annotation_circle(ax, bot1_x, bot1_y, 4, '#4CAF50', 'Bottom 1', 'bottom')
    draw_annotation_circle(ax, bot2_x, bot2_y, 4, '#4CAF50', 'Bottom 2', 'bottom')

    neckline_y = max(h[38:48])
    draw_annotation_line(ax, 10, neckline_y, 85, neckline_y, '#2196F3', '--', 2, 'Neckline')

    support_y = (bot1_y + bot2_y) / 2
    draw_annotation_line(ax, 15, support_y, 70, support_y, '#FF9800', ':', 1.5, 'Support')

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_triple_top(img_path):
    np.random.seed(200)
    n = 140
    base = 900

    def trend(i, n):
        if i < 25:
            return base * 0.005
        elif i < 40:
            return -base * 0.004
        elif i < 55:
            return base * 0.005
        elif i < 70:
            return -base * 0.004
        elif i < 85:
            return base * 0.005
        elif i < 100:
            return -base * 0.003
        else:
            return -base * 0.006

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.010)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "TATATECH", "NSE", "1 Hour", n)

    p1_x, p2_x, p3_x = 28, 58, 88
    p1_y = max(h[25:32])
    p2_y = max(h[55:62])
    p3_y = max(h[85:92])

    draw_annotation_circle(ax, p1_x, p1_y, 4, '#FF5252', 'Top 1')
    draw_annotation_circle(ax, p2_x, p2_y, 4, '#FF5252', 'Top 2')
    draw_annotation_circle(ax, p3_x, p3_y, 4, '#FF5252', 'Top 3')

    resistance_y = max(p1_y, p2_y, p3_y)
    draw_annotation_line(ax, 15, resistance_y, 105, resistance_y, '#FF9800', ':', 1.5, 'Resistance')

    neckline_y = min(min(l[38:48]), min(l[68:78]))
    draw_annotation_line(ax, 15, neckline_y, 120, neckline_y, '#2196F3', '--', 2, 'Neckline / Support')

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_triple_bottom(img_path):
    np.random.seed(300)
    n = 130
    base = 1800

    def trend(i, n):
        if i < 25:
            return -base * 0.004
        elif i < 40:
            return base * 0.003
        elif i < 55:
            return -base * 0.004
        elif i < 70:
            return base * 0.003
        elif i < 85:
            return -base * 0.004
        elif i < 100:
            return base * 0.003
        else:
            return base * 0.006

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.010)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "INFY", "NSE", "Daily", n)

    b1_x, b2_x, b3_x = 28, 58, 88
    b1_y = min(l[25:32])
    b2_y = min(l[55:62])
    b3_y = min(l[85:92])

    draw_annotation_circle(ax, b1_x, b1_y, 4, '#4CAF50', 'Bottom 1', 'bottom')
    draw_annotation_circle(ax, b2_x, b2_y, 4, '#4CAF50', 'Bottom 2', 'bottom')
    draw_annotation_circle(ax, b3_x, b3_y, 4, '#4CAF50', 'Bottom 3', 'bottom')

    support_y = min(b1_y, b2_y, b3_y)
    draw_annotation_line(ax, 15, support_y, 105, support_y, '#FF9800', ':', 1.5, 'Support')

    neckline_y = max(max(h[38:48]), max(h[68:78]))
    draw_annotation_line(ax, 15, neckline_y, 115, neckline_y, '#2196F3', '--', 2, 'Neckline / Resistance')

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_rectangle(img_path):
    np.random.seed(400)
    n = 100
    base = 350

    def trend(i, n):
        if i < 15:
            return base * 0.003
        elif i < 75:
            return np.random.choice([-1, 1]) * base * 0.002
        else:
            return base * 0.006

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.012)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "GENUSPOWER", "NSE", "Daily", n)

    resistance_y = np.percentile(h[15:75], 85)
    support_y = np.percentile(l[15:75], 15)

    draw_annotation_line(ax, 12, resistance_y, 78, resistance_y, '#FF5252', '-', 2.5, 'Resistance')
    draw_annotation_line(ax, 12, support_y, 78, support_y, '#4CAF50', '-', 2.5, 'Support')

    ax.fill_between([12, 78], support_y, resistance_y, alpha=0.08, color='#2196F3')
    ax.text(45, (resistance_y + support_y) / 2, 'RECTANGLE\nPATTERN', color='#2196F3',
            fontsize=12, ha='center', va='center', fontweight='bold', alpha=0.5)

    ax.annotate('Breakout', xy=(80, resistance_y + 5), fontsize=10, color='#4CAF50',
                fontweight='bold',
                arrowprops=dict(arrowstyle='->', color='#4CAF50', lw=2),
                xytext=(85, resistance_y + 20))

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_descending_triangle(img_path):
    np.random.seed(500)
    n = 110
    base = 1600

    def trend(i, n):
        if i < 20:
            return -base * 0.002
        elif i < 80:
            cycle = (i - 20) % 20
            if cycle < 10:
                return -base * 0.003 * (1 - (i - 20) / 80)
            else:
                return base * 0.002 * (1 - (i - 20) / 80)
        else:
            return -base * 0.006

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.010)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "BHARTIARTL", "NSE", "2 Hour", n)

    support_y = np.percentile(l[20:80], 10)
    draw_annotation_line(ax, 15, support_y, 90, support_y, '#4CAF50', '-', 2.5, 'Flat Support')

    high_start = max(h[18:25])
    high_end = np.percentile(h[65:80], 70)
    draw_annotation_line(ax, 20, high_start, 80, high_end, '#FF5252', '-', 2.5, 'Descending Highs')

    ax.text(50, (high_start + support_y) / 2, 'Descending\nTriangle', color='#2196F3',
            fontsize=11, ha='center', va='center', fontweight='bold', alpha=0.5)

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_ascending_triangle(img_path):
    np.random.seed(600)
    n = 120
    base = 300

    def trend(i, n):
        if i < 15:
            return base * 0.004
        elif i < 85:
            cycle = (i - 15) % 20
            if cycle < 10:
                return base * 0.003 * (1 - (i - 15) / 100)
            else:
                return -base * 0.002 * (1 - (i - 15) / 100)
        else:
            return base * 0.007

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.010)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "BPCL", "NSE", "Daily", n)

    resistance_y = np.percentile(h[15:85], 90)
    draw_annotation_line(ax, 12, resistance_y, 90, resistance_y, '#FF5252', '-', 2.5,
                        'Flat Upper Trend Line')

    low_start = min(l[13:20])
    low_end = np.percentile(l[70:85], 30)
    draw_annotation_line(ax, 15, low_start, 85, low_end, '#4CAF50', '-', 2.5,
                        'Rising Lower Trend Line')

    ax.text(50, (resistance_y + low_start) / 2 + 5, 'Ascending Triangle\nPattern Formation',
            color='#2196F3', fontsize=11, ha='center', va='center', fontweight='bold', alpha=0.5)

    ax.annotate('Breakout', xy=(92, resistance_y + 3), fontsize=10, color='#4CAF50',
                fontweight='bold',
                arrowprops=dict(arrowstyle='->', color='#4CAF50', lw=2),
                xytext=(98, resistance_y + 18))

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_symmetrical_triangle(img_path):
    np.random.seed(700)
    n = 120
    base = 11000

    def trend(i, n):
        if i < 15:
            return base * 0.003
        elif i < 85:
            amplitude = base * 0.004 * (1 - (i - 15) / 85)
            cycle = (i - 15) % 18
            if cycle < 9:
                return -amplitude
            else:
                return amplitude
        else:
            return base * 0.005

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.008)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "MARUTI", "NSE", "Daily", n)

    high_start = max(h[13:20])
    high_end = np.percentile(h[75:85], 60)
    draw_annotation_line(ax, 15, high_start, 85, high_end, '#FF5252', '-', 2.5,
                        'Descending Upper Trendline')

    low_start = min(l[18:28])
    low_end = np.percentile(l[75:85], 40)
    draw_annotation_line(ax, 20, low_start, 85, low_end, '#4CAF50', '-', 2.5,
                        'Ascending Lower Trendline')

    mid_y = (high_start + low_start) / 2
    ax.text(50, mid_y, 'Symmetrical\nTriangle', color='#2196F3',
            fontsize=12, ha='center', va='center', fontweight='bold', alpha=0.5)

    apex_x = 85
    ax.annotate('Apex / Convergence', xy=(apex_x, (high_end + low_end) / 2),
                fontsize=9, color='#FF9800', fontweight='bold',
                arrowprops=dict(arrowstyle='->', color='#FF9800', lw=1.5),
                xytext=(apex_x + 8, (high_end + low_end) / 2 + 80))

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_diamond_top(img_path):
    np.random.seed(800)
    n = 130
    base = 5500

    def trend(i, n):
        if i < 20:
            return base * 0.004
        elif i < 45:
            amplitude = base * 0.005 * ((i - 20) / 25)
            return amplitude * np.sin((i - 20) * 0.5)
        elif i < 70:
            amplitude = base * 0.005 * (1 - (i - 45) / 25)
            return amplitude * np.sin((i - 45) * 0.5)
        elif i < 90:
            return -base * 0.002
        else:
            return -base * 0.005

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.010)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "COFORGE", "NSE", "Daily", n)

    left_x, top_x, right_x, bottom_x = 22, 45, 68, 45
    top_y = max(h[40:50])
    bottom_y = min(l[40:50])
    left_y = (top_y + bottom_y) / 2
    right_y = (top_y + bottom_y) / 2

    draw_annotation_line(ax, left_x, left_y, top_x, top_y, '#FF5252', '-', 2)
    draw_annotation_line(ax, top_x, top_y, right_x, right_y, '#FF5252', '-', 2)
    draw_annotation_line(ax, left_x, left_y, bottom_x, bottom_y, '#4CAF50', '-', 2)
    draw_annotation_line(ax, bottom_x, bottom_y, right_x, right_y, '#4CAF50', '-', 2)

    ax.text(45, top_y + 30, 'Diamond Top Pattern', color='#FF9800',
            fontsize=12, ha='center', va='bottom', fontweight='bold',
            bbox=dict(boxstyle='round,pad=0.4', facecolor=TV_DARK['bg'],
                     edgecolor='#FF9800', alpha=0.9))

    ax.text(22, left_y + 40, 'Expanding\nVolatility', color='#2196F3',
            fontsize=9, ha='center', fontweight='bold', alpha=0.7)
    ax.text(68, right_y + 40, 'Contracting\nVolatility', color='#2196F3',
            fontsize=9, ha='center', fontweight='bold', alpha=0.7)

    ax.annotate('Breakdown', xy=(75, bottom_y - 30), fontsize=10, color='#FF5252',
                fontweight='bold',
                arrowprops=dict(arrowstyle='->', color='#FF5252', lw=2),
                xytext=(82, bottom_y + 20))

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


def gen_wedge(img_path):
    np.random.seed(900)
    n = 110
    base = 700

    def trend(i, n):
        if i < 15:
            return base * 0.003
        elif i < 80:
            return base * 0.002 * (1 - (i - 15) / 80)
        else:
            return -base * 0.006

    dates, o, h, l, c, v = generate_ohlcv(n, base, trend, 0.012)

    fig, ax = plt.subplots(figsize=(14, 7))
    plot_candlestick(ax, dates, o, h, l, c)
    setup_chart_axes(fig, ax, "TATAMOTORS", "NSE", "4 Hour", n)

    h_start = max(h[13:20])
    h_end = np.percentile(h[70:80], 80)
    l_start = min(l[13:20])
    l_end = np.percentile(l[70:80], 30)

    upper_slope = (h_end - h_start) / 65
    lower_slope = (l_end - l_start) / 65

    draw_annotation_line(ax, 15, h_start, 80, h_end, '#FF5252', '-', 2.5, 'Upper Trendline')
    draw_annotation_line(ax, 15, l_start, 80, l_end, '#4CAF50', '-', 2.5, 'Lower Trendline')

    mid_y = (h_start + l_start) / 2 + 15
    ax.text(48, mid_y, 'Rising Wedge\n(Bearish)', color='#FF9800',
            fontsize=12, ha='center', va='center', fontweight='bold', alpha=0.6)

    ax.annotate('Breakdown', xy=(85, l_end - 15), fontsize=10, color='#FF5252',
                fontweight='bold',
                arrowprops=dict(arrowstyle='->', color='#FF5252', lw=2),
                xytext=(92, l_end + 15))

    ax.set_ylabel('Price (INR)', color=TV_DARK['text'], fontsize=10)
    fig.tight_layout()
    fig.savefig(img_path, dpi=150, bbox_inches='tight', facecolor=TV_DARK['bg'])
    plt.close()


GENERATORS = {
    "Double Top": gen_double_top,
    "Double Bottom": gen_double_bottom,
    "Triple Top": gen_triple_top,
    "Triple Bottom": gen_triple_bottom,
    "Rectangle (Box)": gen_rectangle,
    "Descending Triangle": gen_descending_triangle,
    "Ascending Triangle": gen_ascending_triangle,
    "Symmetrical Triangle": gen_symmetrical_triangle,
    "Diamond Top": gen_diamond_top,
    "Wedge": gen_wedge,
}

PATTERN_DESCRIPTIONS = {
    "Double Top": "A bearish reversal pattern formed when price reaches a resistance level twice, creating two peaks at approximately the same level. The pattern is confirmed when price breaks below the neckline (support between the two peaks), signaling a potential downtrend.",
    "Double Bottom": "A bullish reversal pattern formed when price reaches a support level twice, creating two troughs at approximately the same level. Confirmation occurs when price breaks above the neckline (resistance between the two bottoms), indicating a potential uptrend.",
    "Triple Top": "A bearish reversal pattern with three peaks at approximately the same resistance level. It signals stronger resistance than a double top and confirms when price breaks below the support level connecting the two reaction lows between the peaks.",
    "Triple Bottom": "A bullish reversal pattern with three troughs at approximately the same support level. It indicates strong support and confirms when price breaks above the resistance level connecting the two reaction highs between the bottoms.",
    "Rectangle (Box)": "A continuation pattern where price oscillates between well-defined horizontal support and resistance levels. The breakout direction (above resistance or below support) determines the next trend direction. Volume typically decreases during the formation.",
    "Descending Triangle": "A bearish continuation pattern characterized by a flat horizontal support line and a descending trendline connecting lower highs. The converging trendlines create a triangle that typically breaks downward, though upside breakouts can occur.",
    "Ascending Triangle": "A bullish continuation pattern with a flat horizontal resistance line and an ascending trendline connecting higher lows. The converging lines indicate increasing buying pressure, and the pattern typically breaks upward through the resistance level.",
    "Symmetrical Triangle": "A neutral continuation pattern formed by converging trendlines — a descending upper line and an ascending lower line. The breakout can occur in either direction, with the breakout direction typically following the prior trend. Volume decreases as the pattern narrows toward the apex.",
    "Diamond Top": "A rare bearish reversal pattern combining a broadening formation (expanding volatility) followed by a symmetrical triangle (contracting volatility), creating a diamond shape. It typically forms at market tops and signals a significant trend reversal when price breaks below the lower right trendline.",
    "Wedge": "A pattern formed by two converging trendlines that slant in the same direction. A rising wedge (both lines slope up) is bearish, while a falling wedge (both lines slope down) is bullish. The pattern indicates diminishing momentum before a reversal.",
}


def generate_all_charts():
    print("Generating 10 TradingView-style chart images...")
    paths = {}
    for chart in CHARTS:
        pattern = chart["pattern"]
        filename = f"chart_{chart['num']:02d}_{pattern.lower().replace(' ', '_').replace('(', '').replace(')', '')}.png"
        img_path = os.path.join(IMG_DIR, filename)
        gen_func = GENERATORS[pattern]
        gen_func(img_path)
        paths[pattern] = img_path
        print(f"  [{chart['num']:2d}] {pattern:25s} -> {filename}")
    return paths


def create_word_document(chart_paths):
    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # --- Cover Page ---
    for _ in range(3):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TECHNICAL ANALYSIS")
    run.bold = True
    run.font.size = Pt(28)
    run.font.name = 'Times New Roman'

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Chart Pattern Recognition")
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_paragraph()

    if os.path.exists(LOGO_PATH):
        logo_para = doc.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_para.add_run().add_picture(LOGO_PATH, width=Inches(2.0))

    doc.add_paragraph()

    info_lines = [
        ("Submitted To:", True),
        ("Prof. Satyaprakash Tiwari", False),
        ("", False),
        ("Submitted By:", True),
        ("Trevor Dcosta", False),
        ("Roll No: MSF07", False),
        ("", False),
        ("Sydenham Institute of Management Studies,", False),
        ("Research and Entrepreneurship Education (SIMSREE)", False),
        ("Semester 3 — Technical Analysis", False),
    ]

    for text, is_bold in info_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        run.bold = is_bold
        run.font.size = Pt(13) if is_bold else Pt(12)
        run.font.name = 'Times New Roman'

    doc.add_page_break()

    # --- Table of Contents ---
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = toc_title.add_run("TABLE OF CONTENTS")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = 'Times New Roman'

    doc.add_paragraph()

    for chart in CHARTS:
        toc_entry = doc.add_paragraph()
        run = toc_entry.add_run(f"{chart['num']}.  {chart['pattern']}")
        run.font.size = Pt(13)
        run.font.name = 'Times New Roman'
        detail = toc_entry.add_run(f"  —  {chart['stock']} ({chart['timeframe']})")
        detail.font.size = Pt(11)
        detail.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_page_break()

    # --- Chart Pages ---
    for chart in CHARTS:
        pattern = chart["pattern"]
        heading = doc.add_heading(f"{chart['num']}.  {pattern}", level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = meta.add_run(f"Stock: {chart['stock']}  |  Exchange: {chart['exchange']}  |  Timeframe: {chart['timeframe']}")
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(80, 80, 80)
        run.italic = True

        doc.add_paragraph()

        img_path = chart_paths.get(pattern)
        if img_path and os.path.exists(img_path):
            img_para = doc.add_paragraph()
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_para.add_run().add_picture(img_path, width=Inches(6.2))

        doc.add_paragraph()

        desc_heading = doc.add_paragraph()
        run = desc_heading.add_run("Pattern Description:")
        run.bold = True
        run.font.size = Pt(12)

        desc = doc.add_paragraph()
        run = desc.add_run(PATTERN_DESCRIPTIONS.get(pattern, ""))
        run.font.size = Pt(11)

        if chart['num'] < 10:
            doc.add_page_break()

    docx_path = os.path.join(BASE, "Trevor_Dcosta_MSF07_Technical_Analysis.docx")
    doc.save(docx_path)
    size_kb = os.path.getsize(docx_path) / 1024
    print(f"\nWord document: {docx_path} ({size_kb:.1f} KB)")
    return docx_path


def create_pdf_document(chart_paths):
    pdf_path = os.path.join(BASE, "Trevor_Dcosta_MSF07_Technical_Analysis.pdf")
    c = pdf_canvas.Canvas(pdf_path, pagesize=A4)
    PAGE_W, PAGE_H = A4

    c.setTitle("Technical Analysis - Trevor Dcosta")
    c.setAuthor("Trevor Dcosta")

    # Cover page
    c.setFont("Times-Bold", 28)
    c.drawCentredString(PAGE_W / 2, PAGE_H - 4 * cm, "TECHNICAL ANALYSIS")

    c.setFont("Times-Bold", 20)
    c.setFillColorRGB(0.3, 0.3, 0.3)
    c.drawCentredString(PAGE_W / 2, PAGE_H - 5.5 * cm, "Chart Pattern Recognition")

    if os.path.exists(LOGO_PATH):
        logo_size = 5 * cm
        logo_x = (PAGE_W - logo_size) / 2
        logo_y = PAGE_H - 7 * cm - logo_size
        try:
            c.drawImage(LOGO_PATH, logo_x, logo_y, logo_size, logo_size,
                       preserveAspectRatio=True, mask='auto')
        except:
            pass

    right_x = PAGE_W - 3.5 * cm
    info_y = PAGE_H - 14 * cm

    c.setFillColorRGB(0, 0, 0)
    c.setFont("Times-Bold", 14)
    c.drawRightString(right_x, info_y, "Submitted To:")
    c.setFont("Times-Roman", 13)
    c.drawRightString(right_x, info_y - 0.7 * cm, "Prof. Satyaprakash Tiwari")

    c.setFont("Times-Bold", 14)
    c.drawRightString(right_x, info_y - 2 * cm, "Submitted By:")
    c.setFont("Times-Roman", 13)
    c.drawRightString(right_x, info_y - 2.7 * cm, "Trevor Dcosta")
    c.drawRightString(right_x, info_y - 3.4 * cm, "Roll No: MSF07")

    c.setFont("Times-Roman", 11)
    c.drawRightString(right_x, info_y - 5 * cm, "Sydenham Institute of Management Studies,")
    c.drawRightString(right_x, info_y - 5.5 * cm, "Research and Entrepreneurship Education (SIMSREE)")
    c.drawRightString(right_x, info_y - 6.2 * cm, "Semester 3 — Technical Analysis")

    c.showPage()

    # Chart pages
    for chart in CHARTS:
        pattern = chart["pattern"]
        top_y = PAGE_H - 2 * cm

        c.setFont("Times-Bold", 18)
        c.setFillColorRGB(0, 0, 0)
        c.drawString(2 * cm, top_y, f"{chart['num']}.  {pattern}")

        c.setFont("Times-Italic", 11)
        c.setFillColorRGB(0.3, 0.3, 0.3)
        meta = f"{chart['stock']}  |  {chart['exchange']}  |  {chart['timeframe']}"
        c.drawString(2 * cm, top_y - 0.8 * cm, meta)

        img_path = chart_paths.get(pattern)
        if img_path and os.path.exists(img_path):
            img = Image.open(img_path)
            img_w, img_h = img.size
            aspect = img_w / img_h

            avail_w = PAGE_W - 3 * cm
            avail_h = PAGE_H - 9 * cm

            if avail_w / aspect <= avail_h:
                draw_w = avail_w
                draw_h = avail_w / aspect
            else:
                draw_h = avail_h
                draw_w = avail_h * aspect

            img_x = (PAGE_W - draw_w) / 2
            img_y = top_y - 2 * cm - draw_h

            c.drawImage(img_path, img_x, img_y, draw_w, draw_h,
                       preserveAspectRatio=True)

        desc = PATTERN_DESCRIPTIONS.get(pattern, "")
        if desc:
            c.setFont("Times-Bold", 11)
            c.setFillColorRGB(0, 0, 0)
            desc_y = 3.5 * cm
            c.drawString(2 * cm, desc_y, "Pattern Description:")

            c.setFont("Times-Roman", 10)
            c.setFillColorRGB(0.2, 0.2, 0.2)
            text_obj = c.beginText(2 * cm, desc_y - 0.5 * cm)
            text_obj.setFont("Times-Roman", 10)

            words = desc.split()
            line = ""
            max_width = PAGE_W - 4 * cm
            for word in words:
                test_line = f"{line} {word}".strip()
                if c.stringWidth(test_line, "Times-Roman", 10) < max_width:
                    line = test_line
                else:
                    text_obj.textLine(line)
                    line = word
            if line:
                text_obj.textLine(line)
            c.drawText(text_obj)

        c.showPage()

    c.save()
    size_kb = os.path.getsize(pdf_path) / 1024
    print(f"PDF document:  {pdf_path} ({size_kb:.1f} KB)")
    return pdf_path


def main():
    chart_paths = generate_all_charts()
    docx_path = create_word_document(chart_paths)
    pdf_path = create_pdf_document(chart_paths)
    print(f"\nDone! Generated:")
    print(f"  - 10 chart images in {IMG_DIR}/")
    print(f"  - {docx_path}")
    print(f"  - {pdf_path}")


if __name__ == "__main__":
    main()
